"""Aggiorna l'Excel dei codici Mexal dall'esportazione di Mexal e controlla products.json.

Uso (dalla cartella CiliegioShop):
    python aggiorna_codici_mexal.py

Legge:   Export prodotti mexal.xlsx   (esportato da Mexal: Codice, Descrizione, Um, Cod.alternativo, Natura, Gruppo)
         products.json                (campo "mexal" di ogni prodotto dello shop)
Scrive:  ../../Codici Mexal Siena Wine.xlsx (cartella LUCA)   (foglio "Codici Mexal" = copia dell'esportazione,
                                            foglio "Abbinamento shop" = prodotti shop + codice Mexal)
Segnala: prodotti dello shop senza codice, o con un codice che non esiste piu' in Mexal;
         codici Mexal presenti nell'esportazione ma non nello shop (solo conteggio).
Il foglio "Abbinamento shop" ha la colonna "IVA % (da compilare)": l'aliquota IVA di ogni prodotto.
Si compila a mano nell'Excel e poi si lancia:
    python aggiorna_codici_mexal.py --importa-iva
che copia le aliquote in products.json (campo "vat"). Quando si rigenera l'Excel le aliquote gia' scritte
(in products.json o nel vecchio Excel) vengono mantenute.
Non modifica altro di products.json: i codici Mexal si assegnano a mano.
"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
EXPORT = os.path.join(HERE, "Export prodotti mexal.xlsx")
PRODUCTS = os.path.join(HERE, "products.json")
OUT = os.path.join(os.path.dirname(os.path.dirname(HERE)), "Codici Mexal Siena Wine.xlsx")


def leggi_export():
    wb = openpyxl.load_workbook(EXPORT, read_only=True)
    ws = wb.worksheets[0]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    # la prima colonna dell'export e' vuota: l'intestazione "Codice" e' nella seconda
    start = next(i for i, r in enumerate(rows) if r and "Codice" in [str(x) for x in r if x])
    head = [str(x) if x is not None else "" for x in rows[start]]
    ci = head.index("Codice")
    out = []
    for r in rows[start + 1:]:
        if r[ci] is None:
            continue
        g = lambda name: ("" if r[head.index(name)] is None else str(r[head.index(name)]).strip())
        out.append([str(r[ci]).strip(), g("Descrizione"), g("Um"), g("Cod.alternativo"), g("Natura"),
                    g("Gruppo merceologico")])
    return out


def leggi_iva_da_excel():
    """Aliquote IVA scritte nel foglio 'Abbinamento shop' dell'Excel esistente: {id shop: numero}."""
    out = {}
    if not os.path.exists(OUT):
        return out
    wb = openpyxl.load_workbook(OUT, read_only=True)
    if "Abbinamento shop" not in wb.sheetnames:
        return out
    ws = wb["Abbinamento shop"]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    if not rows or "IVA % (da compilare)" not in rows[0]:
        return out
    ci = list(rows[0]).index("IVA % (da compilare)")
    for r in rows[1:]:
        if r[0] is None or len(r) <= ci or r[ci] in (None, ""):
            continue
        try:
            out[str(r[0]).strip()] = float(str(r[ci]).replace(",", ".").replace("%", "").strip())
        except ValueError:
            print(f"ATTENZIONE: aliquota non valida per id {r[0]}: {r[ci]!r} (ignorata)")
    return out


def importa_iva():
    iva = leggi_iva_da_excel()
    if not iva:
        sys.exit("Nessuna aliquota trovata nella colonna 'IVA % (da compilare)' di " + OUT)
    prods = json.load(open(PRODUCTS, encoding="utf-8"))
    n = 0
    for p in prods:
        pid = str(p["id"])
        v = iva.get(pid)
        if v is None and pid.endswith("_gift"):
            v = iva.get(pid[:-5])
        if v is None:
            continue
        if not (0 <= v <= 30):
            sys.exit(f"Aliquota fuori range per id {pid}: {v}")
        p["vat"] = int(v) if float(v).is_integer() else v
        n += 1
    with open(PRODUCTS, "w", encoding="utf-8", newline="\n") as f:
        json.dump(prods, f, ensure_ascii=False, indent=2)
    print(f"Aliquote IVA copiate in products.json per {n} prodotti su {len(prods)}.")
    senza = [f"{p['id']} {p['name']}" for p in prods if "vat" not in p]
    if senza:
        print("Ancora senza aliquota:")
        for x in senza:
            print(" -", x)


def main():
    if "--importa-iva" in sys.argv:
        importa_iva()
        return
    if not os.path.exists(EXPORT):
        sys.exit("Manca il file: " + EXPORT)
    mexal = leggi_export()
    by_code = {r[0]: r for r in mexal}
    prods = json.load(open(PRODUCTS, encoding="utf-8"))

    problemi = []
    for p in prods:
        c = p.get("mexal")
        if not c:
            problemi.append(f"id {p['id']} ({p['name']}): nessun codice Mexal")
        elif c not in by_code:
            problemi.append(f"id {p['id']} ({p['name']}): il codice {c} non esiste nell'esportazione Mexal")
    usati = {p.get("mexal") for p in prods}

    iva_old = leggi_iva_da_excel()
    wb = openpyxl.Workbook()
    hdr = PatternFill("solid", fgColor="D4AF37")

    def foglio(ws, head, rows, widths):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = hdr
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.number_format = "@"

    ws = wb.active
    ws.title = "Codici Mexal"
    foglio(ws, ["Codice", "Descrizione", "UM", "Cod. alternativo (EAN)", "Natura", "Gruppo merceologico", "Nello shop"],
           [r + ["si" if r[0] in usati else ""] for r in mexal], [12, 70, 6, 22, 9, 12, 10])

    ws2 = wb.create_sheet("Abbinamento shop")
    rows = []
    for p in prods:
        c = p.get("mexal", "")
        iva = p.get("vat", iva_old.get(str(p["id"]), iva_old.get(str(p["id"]).replace("_gift", ""), "")))
        rows.append([str(p["id"]), p["name"], p["category"], "si" if p.get("available", True) else "no",
                     float(p["price"]), c, by_code[c][1] if c in by_code else "NON TROVATO", iva])
    foglio(ws2, ["ID shop", "Nome nello shop", "Categoria", "Disponibile", "Prezzo EUR", "Codice Mexal",
                 "Descrizione Mexal", "IVA % (da compilare)"], rows, [9, 58, 13, 12, 11, 13, 70, 20])
    for row in ws2.iter_rows(min_row=2):
        row[4].number_format = "0.00"
        row[7].number_format = "General"
        row[7].fill = PatternFill("solid", fgColor="FFF2CC")
    wb.save(OUT)

    print(f"Codici Mexal letti: {len(mexal)} | prodotti shop: {len(prods)} | non nello shop: {len(mexal) - len(usati & set(by_code))}")
    print("Scritto:", OUT)
    if problemi:
        print("\nATTENZIONE:")
        for x in problemi:
            print(" -", x)
        sys.exit(1)
    print("Tutti i prodotti dello shop hanno un codice Mexal valido.")


if __name__ == "__main__":
    main()
